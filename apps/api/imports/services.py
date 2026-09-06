import hashlib
from decimal import Decimal
from io import BytesIO

from django.db import transaction
from django.utils import timezone
from django.utils.dateparse import parse_date
from django.utils.text import slugify
from openpyxl import Workbook, load_workbook

from audit.models import record_audit
from operations.models import Client, Driver, Indent, Trip, TripCharge, Vehicle, Vendor

from .models import ImportJob
from .parsing import parse_date as parse_cell_date
from .parsing import parse_decimal

LEGACY_COLUMNS = [
    "SR NO",
    "FROM",
    "TO",
    "DATE",
    "UOM-LTRS",
    "QTY",
    "TOTAL LOAD",
    "RATES FOR VEHICLES",
    "UNLOADING",
    "ADVANCE",
    "TOTAL PAYMENTS TO BE DONE",
    "TRANSPORTER",
    "VH NO",
    "DRIVER NAME",
    "DRIVER NO",
    "VH TYPE",
    "EDD",
]


def _decimal(value, field, errors):
    return parse_decimal(value, field, errors, required=False, default=Decimal("0.00"))


def _date(value, field, errors, epoch):
    return parse_cell_date(value, field, errors, epoch)


def preview_legacy_workbook(content):
    formula_book = load_workbook(BytesIO(content), read_only=True, data_only=False)
    value_book = load_workbook(BytesIO(content), read_only=True, data_only=True)
    formula_sheet = formula_book.active
    value_sheet = value_book.active
    headers = [str(cell.value).strip().upper() if cell.value is not None else "" for cell in next(formula_sheet.iter_rows(min_row=1, max_row=1))]
    missing = [column for column in LEGACY_COLUMNS if column not in headers]
    if missing:
        raise ValueError(f"Missing required columns: {', '.join(missing)}")
    positions = {name: headers.index(name) for name in LEGACY_COLUMNS}
    formula_rows = formula_sheet.iter_rows(min_row=2, values_only=True)
    value_rows = value_sheet.iter_rows(min_row=2, values_only=True)
    rows = []
    fingerprints = set()
    for row_no, (formula_values, cached_values) in enumerate(zip(formula_rows, value_rows, strict=False), start=2):
        if not any(value not in (None, "") for value in formula_values):
            continue
        if len(rows) >= 1000:
            raise ValueError("Legacy imports are limited to 1,000 populated rows per workbook")
        def get_formula(name, values=formula_values):
            return values[positions[name]]

        def get_value(name, values=cached_values):
            return values[positions[name]]

        errors, warnings = [], []
        origin = str(get_formula("FROM") or "").strip()
        destination = str(get_formula("TO") or "").strip()
        vendor = str(get_formula("TRANSPORTER") or "").strip()
        vehicle = "".join(ch for ch in str(get_formula("VH NO") or "").upper() if ch.isalnum())
        if not origin:
            errors.append("FROM is required")
        if not destination:
            errors.append("TO is required")
        if not vendor:
            errors.append("TRANSPORTER is required")
        if not vehicle:
            errors.append("VH NO is required")
        deployment_date = _date(get_formula("DATE"), "DATE", errors, formula_book.epoch)
        edd = _date(get_formula("EDD"), "EDD", errors, formula_book.epoch)
        freight = _decimal(get_formula("RATES FOR VEHICLES"), "RATES FOR VEHICLES", errors)
        unloading = _decimal(get_formula("UNLOADING"), "UNLOADING", errors)
        expected_advance = (freight * Decimal("0.90")).quantize(Decimal("0.01")) if freight is not None else None
        expected_total = expected_advance + unloading if expected_advance is not None and unloading is not None else None
        cached_advance = get_value("ADVANCE")
        cached_total = get_value("TOTAL PAYMENTS TO BE DONE")
        supplied_advance = None if cached_advance in (None, "") else _decimal(cached_advance, "ADVANCE", [])
        supplied_total = None if cached_total in (None, "") else _decimal(cached_total, "TOTAL PAYMENTS TO BE DONE", [])
        if supplied_advance is not None and expected_advance is not None and supplied_advance != expected_advance:
            warnings.append(f"ADVANCE differs from calculated 90% ({expected_advance})")
        if supplied_total is not None and expected_total is not None and supplied_total != expected_total:
            warnings.append(f"TOTAL PAYMENTS differs from advance + unloading ({expected_total})")
        fingerprint = (deployment_date, origin.lower(), destination.lower(), vendor.lower(), vehicle, freight)
        if fingerprint in fingerprints:
            warnings.append("Possible duplicate row in this workbook")
        fingerprints.add(fingerprint)
        rows.append(
            {
                "row_no": row_no,
                "valid": not errors,
                "errors": errors,
                "warnings": warnings,
                "normalized": {
                    "source_row_no": get_formula("SR NO"),
                    "origin": origin,
                    "destination": destination,
                    "deployment_date": deployment_date.isoformat() if deployment_date else None,
                    "uom_ltrs": str(get_formula("UOM-LTRS") or ""),
                    "quantity": str(get_formula("QTY") or ""),
                    "total_load": str(get_formula("TOTAL LOAD") or ""),
                    "vendor_freight_rate": str(freight) if freight is not None else None,
                    "unloading": str(unloading) if unloading is not None else None,
                    "calculated_advance": str(expected_advance) if expected_advance is not None else None,
                    "calculated_total": str(expected_total) if expected_total is not None else None,
                    "vendor": vendor,
                    "vehicle_registration": vehicle,
                    "driver_name": str(get_formula("DRIVER NAME") or "").strip(),
                    "driver_phone": str(get_formula("DRIVER NO") or "").strip(),
                    "vehicle_type": str(get_formula("VH TYPE") or "").strip(),
                    "expected_delivery_date": edd.isoformat() if edd else None,
                },
            }
        )
    return {
        "source_hash": hashlib.sha256(content).hexdigest(),
        "columns": headers,
        "required_columns": LEGACY_COLUMNS,
        "row_count": len(rows),
        "valid_count": sum(row["valid"] for row in rows),
        "error_count": sum(not row["valid"] for row in rows),
        "rows": rows,
    }


def _vendor_code(name):
    base = slugify(name).replace("-", "").upper()[:20] or "VENDOR"
    candidate = base
    suffix = 1
    while Vendor.objects.filter(vendor_code=candidate).exists():
        suffix += 1
        candidate = f"{base[:16]}{suffix:04d}"
    return candidate


@transaction.atomic
def confirm_import(
    *, job, actor, client_id, create_missing=False, allow_partial=False, import_duplicates=False, request_id=""
):
    job = ImportJob.objects.select_for_update().get(pk=job.pk)
    if job.status == "COMPLETED":
        return job
    if job.uploaded_by_id != actor.pk and getattr(actor, "role", "") != "ADMIN":
        raise PermissionError("Only the uploader or an administrator can confirm this import")
    client = Client.objects.get(pk=client_id)
    rows = job.summary.get("rows", [])
    invalid = [row for row in rows if not row.get("valid")]
    if invalid and not allow_partial:
        raise ValueError("The import contains invalid rows; correct them or enable valid-row-only import")
    preflight_errors = []
    for row in rows:
        if not row.get("valid"):
            continue
        values = row["normalized"]
        deployment_date = values["deployment_date"]
        expected_delivery_date = values["expected_delivery_date"]
        if isinstance(deployment_date, str):
            deployment_date = parse_date(deployment_date)
        if isinstance(expected_delivery_date, str):
            expected_delivery_date = parse_date(expected_delivery_date)
        if not deployment_date:
            raise ValueError(f"Row {row['row_no']}: deployment date is invalid")
        vendor_exists = Vendor.objects.filter(display_name__iexact=values["vendor"]).exists()
        if not vendor_exists and not create_missing:
            preflight_errors.append({"row_no": row["row_no"], "error": f"Vendor '{values['vendor']}' does not exist"})
            continue
        vendor = Vendor.objects.filter(display_name__iexact=values["vendor"]).first()
        vehicle_exists = vendor and Vehicle.objects.filter(
            registration_no=values["vehicle_registration"], vendor=vendor
        ).exists()
        if not vehicle_exists and not create_missing:
            preflight_errors.append(
                {"row_no": row["row_no"], "error": f"Vehicle '{values['vehicle_registration']}' does not exist for this vendor"}
            )
            continue
        foreign_vehicle = Vehicle.objects.filter(registration_no=values["vehicle_registration"])
        if vendor:
            foreign_vehicle = foreign_vehicle.exclude(vendor=vendor)
        if foreign_vehicle.exists():
            preflight_errors.append(
                {"row_no": row["row_no"], "error": f"Vehicle '{values['vehicle_registration']}' belongs to another vendor"}
            )
    if preflight_errors:
        job.result = {"errors": preflight_errors, "created_trip_ids": [], "skipped": []}
        job.status = "VALIDATION_FAILED"
        job.save(update_fields=["result", "status", "updated_at"])
        return job

    created, skipped = [], []
    for row in rows:
        if not row.get("valid"):
            skipped.append({"row_no": row["row_no"], "reason": "validation_error"})
            continue
        values = row["normalized"]
        deployment_date = values["deployment_date"]
        expected_delivery_date = values["expected_delivery_date"]
        if isinstance(deployment_date, str):
            deployment_date = parse_date(deployment_date)
        if isinstance(expected_delivery_date, str):
            expected_delivery_date = parse_date(expected_delivery_date)
        vendor = Vendor.objects.filter(display_name__iexact=values["vendor"]).order_by("pk").first()
        if vendor is None:
            vendor = Vendor.objects.create(
                vendor_code=_vendor_code(values["vendor"]),
                display_name=values["vendor"],
                legal_name=values["vendor"],
            )
        duplicate = Trip.objects.filter(
            deployment_date=deployment_date,
            origin__iexact=values["origin"],
            destination__iexact=values["destination"],
            vendor=vendor,
            vehicle_registration_snapshot=values["vehicle_registration"],
            vendor_freight_rate=values["vendor_freight_rate"],
        ).first()
        if duplicate and not import_duplicates:
            skipped.append({"row_no": row["row_no"], "reason": "database_duplicate", "trip_id": duplicate.pk})
            continue
        vehicle, _ = Vehicle.objects.get_or_create(
            registration_no=values["vehicle_registration"],
            defaults={
                "vendor": vendor,
                "vehicle_type": values["vehicle_type"] or "UNSPECIFIED",
            },
        )
        if vehicle.vendor_id != vendor.pk:
            raise ValueError(f"Row {row['row_no']}: vehicle belongs to another vendor")
        driver = Driver.objects.filter(name__iexact=values["driver_name"], phone=values["driver_phone"], vendor=vendor).first()
        if not driver:
            driver = Driver.objects.create(
                name=values["driver_name"] or "Unspecified driver",
                phone=values["driver_phone"],
                vendor=vendor,
            )
        indent_no = f"IMP-{job.pk}-{row['row_no']}"
        indent = Indent.objects.create(
            indent_no=indent_no,
            client=client,
            indent_date=deployment_date,
            origin=values["origin"],
            destination=values["destination"],
            expected_delivery_date=expected_delivery_date,
            uom_ltrs=values["uom_ltrs"],
            quantity=Decimal(values["quantity"] or "0") or None,
            total_load=Decimal(values["total_load"] or "0") or None,
            required_vehicle_type=values["vehicle_type"],
            customer_reference=f"Legacy row {values['source_row_no']}",
        )
        trip = Trip.objects.create(
            indent=indent,
            client=client,
            origin=values["origin"],
            destination=values["destination"],
            deployment_date=deployment_date,
            expected_delivery_date=expected_delivery_date,
            uom_ltrs=values["uom_ltrs"],
            quantity=Decimal(values["quantity"] or "0") or None,
            total_load=Decimal(values["total_load"] or "0") or None,
            vendor=vendor,
            vehicle=vehicle,
            driver=driver,
            vendor_freight_rate=values["vendor_freight_rate"],
        )
        if Decimal(values["unloading"] or "0"):
            TripCharge.objects.create(
                trip=trip,
                charge_type=TripCharge.ChargeType.UNLOADING,
                amount=values["unloading"],
                advance_eligible=True,
                created_by=actor,
            )
        record_audit(
            actor=actor,
            action="TRIP_IMPORTED",
            instance=trip,
            after={"import_job": job.pk, "source_row": row["row_no"]},
            request_id=request_id,
            source="IMPORT",
        )
        created.append(trip.pk)
    job.status = "COMPLETED"
    job.confirmed_at = timezone.now()
    job.result = {"created_trip_ids": created, "skipped": skipped, "errors": []}
    job.save(update_fields=["status", "confirmed_at", "result", "updated_at"])
    record_audit(
        actor=actor,
        action="EXCEL_IMPORT_CONFIRMED",
        instance=job,
        after={"created": len(created), "skipped": len(skipped)},
        request_id=request_id,
        source="IMPORT",
    )
    return job


def build_legacy_export(queryset):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Payment Approval"
    sheet.append(LEGACY_COLUMNS)
    for index, trip in enumerate(queryset.select_related("vendor", "vehicle", "driver").prefetch_related("charges"), start=1):
        unloading = sum(
            (charge.amount for charge in trip.charges.all() if charge.charge_type == TripCharge.ChargeType.UNLOADING),
            Decimal("0"),
        )
        advance = (trip.vendor_freight_rate * trip.advance_percent / Decimal("100")).quantize(Decimal("0.01"))
        sheet.append(
            [
                index,
                trip.origin,
                trip.destination,
                trip.deployment_date,
                trip.uom_ltrs,
                trip.quantity,
                trip.total_load,
                trip.vendor_freight_rate,
                unloading,
                advance,
                advance + unloading,
                trip.vendor.display_name,
                trip.vehicle_registration_snapshot,
                trip.driver_name_snapshot,
                trip.driver_phone_snapshot,
                trip.vehicle_type_snapshot,
                trip.expected_delivery_date,
            ]
        )
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()
