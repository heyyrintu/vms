import hashlib
import re
from collections import defaultdict
from datetime import date
from decimal import Decimal
from io import BytesIO

from django.db import transaction
from django.db.models import Prefetch, Q
from django.utils import timezone
from django.utils.dateparse import parse_date
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

from approvals.models import ApprovalAction, PaymentApprovalBatch, PaymentApprovalItem
from audit.models import record_audit
from core.models import NumberSequence, OrganizationSettings, TDSPolicy
from operations.models import Client, Driver, Indent, Trip, TripCharge, Vehicle, Vendor
from payments.models import FinancePaymentTransaction, PaymentAllocation
from payments.services import create_paid_payment

from .models import ImportJob
from .parsing import parse_text
from .services import _date, _decimal, _vendor_code

MIS_COLUMNS = [
    "RECORD REF",
    "TRIP NO",
    "CLIENT CODE",
    "INDENT NO",
    "FROM",
    "TO",
    "DEPLOYMENT DATE",
    "EDD",
    "TRANSPORTER CODE",
    "TRANSPORTER NAME",
    "VEHICLE NO",
    "VEHICLE TYPE",
    "DRIVER NAME",
    "DRIVER NO",
    "FREIGHT AMOUNT",
    "ADVANCE PERCENT",
    "UNLOADING",
    "GROSS APPROVED",
    "TDS RATE",
    "TDS BASE",
    "TDS APPROVED",
    "NET APPROVED",
    "TDS PAID",
    "NET PAID",
    "PAYMENT DATE",
    "UTR",
    "PAYMENT MODE",
    "TRIP STATUS",
    "NOTES",
]


def _text(value):
    return parse_text(value)


def _money(value, field, errors, default="0.00"):
    amount = _decimal(default if value in (None, "") else value, field, errors)
    if amount is not None and amount < 0:
        errors.append(f"{field} cannot be negative")
    return amount


def _trip_key(values):
    if values["trip_no"]:
        return values["trip_no"].upper()
    if values["record_ref"]:
        return f"REF|{values['client_code'].upper()}|{values['record_ref'].upper()}"
    if values["indent_no"]:
        return f"INDENT|{values['client_code'].upper()}|{values['indent_no'].upper()}|{values['vehicle_registration']}"
    return "|".join(
        [
            values["client_code"].upper(),
            values["deployment_date"] or "",
            values["origin"].lower(),
            values["destination"].lower(),
            values["vendor_name"].lower(),
            values["vehicle_registration"],
        ]
    )


def preview_mis_workbook(content):
    workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
    sheet = workbook["MIS Records"] if "MIS Records" in workbook.sheetnames else workbook.active
    first_row = next(sheet.iter_rows(min_row=1, max_row=1), None)
    if not first_row:
        raise ValueError("The workbook is empty")
    headers = [_text(cell.value).upper() for cell in first_row]
    missing = [column for column in MIS_COLUMNS if column not in headers]
    if missing:
        raise ValueError(f"Missing required columns: {', '.join(missing)}")
    positions = {name: headers.index(name) for name in MIS_COLUMNS}
    settings = OrganizationSettings.load()
    rows = []
    payment_fingerprints = set()
    trip_financials = {}

    for row_no, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
        if not any(value not in (None, "") for value in row):
            continue
        if len(rows) >= 5000:
            raise ValueError("MIS imports are limited to 5,000 rows per workbook")

        def get(name, values=row):
            return values[positions[name]] if positions[name] < len(values) else None

        errors, warnings = [], []
        deployment_date = _date(get("DEPLOYMENT DATE"), "DEPLOYMENT DATE", errors, workbook.epoch)
        expected_delivery_date = _date(get("EDD"), "EDD", errors, workbook.epoch)
        payment_date = _date(get("PAYMENT DATE"), "PAYMENT DATE", errors, workbook.epoch)
        freight = _money(get("FREIGHT AMOUNT"), "FREIGHT AMOUNT", errors)
        advance_percent = _money(
            get("ADVANCE PERCENT"), "ADVANCE PERCENT", errors, str(settings.default_advance_percent)
        )
        unloading = _money(get("UNLOADING"), "UNLOADING", errors)
        tds_rate = _money(get("TDS RATE"), "TDS RATE", errors, str(settings.default_tds_rate))
        if advance_percent is not None and advance_percent > 100:
            errors.append("ADVANCE PERCENT cannot exceed 100")
        if tds_rate is not None and tds_rate > 100:
            errors.append("TDS RATE cannot exceed 100")

        calculated_advance = (
            (freight * advance_percent / Decimal("100")).quantize(Decimal("0.01"))
            if freight is not None and advance_percent is not None
            else None
        )
        calculated_gross = (
            calculated_advance + unloading
            if calculated_advance is not None and unloading is not None
            else None
        )
        gross_approved = _money(
            get("GROSS APPROVED"),
            "GROSS APPROVED",
            errors,
            str(calculated_gross or Decimal("0")),
        )
        tds_base = _money(get("TDS BASE"), "TDS BASE", errors, str(gross_approved or Decimal("0")))
        calculated_tds = (
            (tds_base * tds_rate / Decimal("100")).quantize(Decimal("0.01"))
            if tds_base is not None and tds_rate is not None
            else None
        )
        tds_approved = _money(
            get("TDS APPROVED"), "TDS APPROVED", errors, str(calculated_tds or Decimal("0"))
        )
        expected_net_approved = (
            gross_approved - tds_approved
            if gross_approved is not None and tds_approved is not None
            else None
        )
        net_approved = _money(
            get("NET APPROVED"),
            "NET APPROVED",
            errors,
            str(expected_net_approved or Decimal("0")),
        )
        tds_paid = _money(get("TDS PAID"), "TDS PAID", errors)
        net_paid = _money(get("NET PAID"), "NET PAID", errors)
        gross_paid = (
            tds_paid + net_paid if tds_paid is not None and net_paid is not None else None
        )

        client_code = _text(get("CLIENT CODE")).upper()
        origin = _text(get("FROM"))
        destination = _text(get("TO"))
        vendor_name = _text(get("TRANSPORTER NAME"))
        vehicle_registration = "".join(
            character for character in _text(get("VEHICLE NO")).upper() if character.isalnum()
        )
        trip_status = _text(get("TRIP STATUS")).upper()
        utr = _text(get("UTR"))

        for field, value in (
            ("CLIENT CODE", client_code),
            ("FROM", origin),
            ("TO", destination),
            ("TRANSPORTER NAME", vendor_name),
            ("VEHICLE NO", vehicle_registration),
        ):
            if not value:
                errors.append(f"{field} is required")
        if gross_approved is not None and tds_approved is not None and tds_approved > gross_approved:
            errors.append("TDS APPROVED cannot exceed GROSS APPROVED")
        if expected_net_approved is not None and net_approved != expected_net_approved:
            errors.append("NET APPROVED must equal GROSS APPROVED minus TDS APPROVED")
        if gross_paid is not None and gross_approved is not None and gross_paid > gross_approved:
            errors.append("TDS PAID plus NET PAID cannot exceed GROSS APPROVED")
        if tds_paid is not None and tds_approved is not None and tds_paid > tds_approved:
            errors.append("TDS PAID cannot exceed TDS APPROVED")
        if net_paid is not None and net_approved is not None and net_paid > net_approved:
            errors.append("NET PAID cannot exceed NET APPROVED")
        if gross_paid and (not payment_date or not utr):
            errors.append("PAYMENT DATE and UTR are required when a paid amount is entered")
        if not gross_paid and (payment_date or utr):
            warnings.append("Payment reference was supplied without a paid amount and will not be posted")
        if trip_status and trip_status not in Trip.Status.values:
            errors.append(f"TRIP STATUS must be one of: {', '.join(Trip.Status.values)}")

        values = {
            "record_ref": _text(get("RECORD REF")),
            "trip_no": _text(get("TRIP NO")),
            "client_code": client_code,
            "indent_no": _text(get("INDENT NO")),
            "origin": origin,
            "destination": destination,
            "deployment_date": deployment_date.isoformat() if deployment_date else None,
            "expected_delivery_date": expected_delivery_date.isoformat() if expected_delivery_date else None,
            "vendor_code": _text(get("TRANSPORTER CODE")).upper(),
            "vendor_name": vendor_name,
            "vehicle_registration": vehicle_registration,
            "vehicle_type": _text(get("VEHICLE TYPE")) or "UNSPECIFIED",
            "driver_name": _text(get("DRIVER NAME")) or "Unspecified driver",
            "driver_phone": _text(get("DRIVER NO")),
            "freight_amount": str(freight) if freight is not None else None,
            "advance_percent": str(advance_percent) if advance_percent is not None else None,
            "unloading": str(unloading) if unloading is not None else None,
            "calculated_advance": str(calculated_advance) if calculated_advance is not None else None,
            "gross_approved": str(gross_approved) if gross_approved is not None else None,
            "tds_rate": str(tds_rate) if tds_rate is not None else None,
            "tds_base": str(tds_base) if tds_base is not None else None,
            "tds_approved": str(tds_approved) if tds_approved is not None else None,
            "net_approved": str(net_approved) if net_approved is not None else None,
            "tds_paid": str(tds_paid) if tds_paid is not None else None,
            "net_paid": str(net_paid) if net_paid is not None else None,
            "gross_paid": str(gross_paid) if gross_paid is not None else None,
            "payment_date": payment_date.isoformat() if payment_date else None,
            "utr": utr,
            "payment_mode": _text(get("PAYMENT MODE")).upper() or "BANK_TRANSFER",
            "trip_status": trip_status,
            "notes": _text(get("NOTES")),
        }
        trip_key = _trip_key(values)
        financial_signature = tuple(
            values[field]
            for field in (
                "freight_amount",
                "advance_percent",
                "unloading",
                "gross_approved",
                "tds_rate",
                "tds_base",
                "tds_approved",
                "net_approved",
            )
        )
        if trip_key in trip_financials and trip_financials[trip_key] != financial_signature:
            errors.append("Repeated trip rows must use the same approved financial values")
        trip_financials.setdefault(trip_key, financial_signature)
        if gross_paid:
            payment_fingerprint = (
                values["vendor_name"].lower(),
                values["payment_date"],
                values["utr"].lower(),
                trip_key,
            )
            if payment_fingerprint in payment_fingerprints:
                errors.append("Duplicate trip allocation for the same UTR in this workbook")
            payment_fingerprints.add(payment_fingerprint)
        values["trip_key"] = trip_key
        rows.append(
            {
                "row_no": row_no,
                "valid": not errors,
                "errors": errors,
                "warnings": warnings,
                "normalized": values,
            }
        )

    return {
        "source_hash": hashlib.sha256(b"MIS_HISTORY\0" + content).hexdigest(),
        "columns": headers,
        "required_columns": MIS_COLUMNS,
        "row_count": len(rows),
        "valid_count": sum(row["valid"] for row in rows),
        "error_count": sum(not row["valid"] for row in rows),
        "rows": rows,
    }


def build_mis_template():
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "MIS Records"
    sheet.append(MIS_COLUMNS)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = "A1:AC1"
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="D71920")
        cell.alignment = Alignment(horizontal="center")
    widths = {
        "A": 16, "B": 20, "C": 14, "D": 18, "E": 18, "F": 18, "G": 18, "H": 14,
        "I": 20, "J": 25, "K": 17, "L": 18, "M": 22, "N": 18, "O": 18,
        "P": 17, "Q": 14, "R": 18, "S": 12, "T": 15, "U": 17, "V": 17,
        "W": 14, "X": 14, "Y": 16, "Z": 22, "AA": 18, "AB": 28, "AC": 34,
    }
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width

    example = workbook.create_sheet("Example")
    example.append(MIS_COLUMNS)
    example.append(
        [
            "OLD-0001", "NPL-2024-000001", "NPL", "IND-2024-001", "Sonipat", "Delhi",
            date(2024, 7, 27), date(2024, 7, 28), "V-001", "Sample Transporter",
            "HR10AB1234", "32 FT", "Sample Driver", "9000000000", 50000, 90, 2500,
            47500, 1, 45000, 450, 47050, 450, 47050, date(2024, 7, 27),
            "SAMPLE-UTR-001", "BANK_TRANSFER", "DELIVERED", "Example only — do not upload this sheet",
        ]
    )
    instructions = workbook.create_sheet("Instructions")
    instructions.append(["Drona Logitech MIS historical record upload"])
    instructions.append(["Enter data only on the MIS Records sheet. Do not rename or delete columns."])
    instructions.append(["One row represents one trip/payment allocation. Repeat a trip on another row for another UTR."])
    instructions.append(["Repeated trip rows must have identical freight, approval and TDS values."])
    instructions.append(["If TDS/approval cells are blank, organization defaults and the 90% calculation are used."])
    instructions.append(["PAYMENT DATE and UTR are mandatory when TDS PAID or NET PAID is entered."])
    instructions.append(["Confirming an import is admin-only and creates audited approval/payment ledger entries."])
    instructions.column_dimensions["A"].width = 115
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def _resolve_vendor(values, create_missing):
    vendor = None
    if values["vendor_code"]:
        vendor = Vendor.objects.filter(vendor_code__iexact=values["vendor_code"]).first()
    if not vendor:
        vendor = Vendor.objects.filter(display_name__iexact=values["vendor_name"]).first()
    if vendor:
        return vendor
    if not create_missing:
        raise ValueError(f"Transporter '{values['vendor_name']}' does not exist")
    return Vendor.objects.create(
        vendor_code=values["vendor_code"] or _vendor_code(values["vendor_name"]),
        display_name=values["vendor_name"],
        legal_name=values["vendor_name"],
    )


def _resolve_trip(values, actor, job, create_missing):
    client = Client.objects.filter(code__iexact=values["client_code"]).first()
    if not client:
        raise ValueError(f"Client '{values['client_code']}' does not exist")
    vendor = _resolve_vendor(values, create_missing)
    trip = Trip.objects.filter(trip_no__iexact=values["trip_no"]).first() if values["trip_no"] else None
    if not trip and values["record_ref"]:
        trip = Trip.objects.filter(
            client=client,
            indent__customer_reference__iexact=values["record_ref"],
        ).first()
    if not trip and values["indent_no"]:
        trip = Trip.objects.filter(
            client=client,
            indent__indent_no__iexact=values["indent_no"],
            vehicle_registration_snapshot=values["vehicle_registration"],
        ).first()
    if not trip:
        trip = Trip.objects.filter(
            client=client,
            deployment_date=parse_date(values["deployment_date"]),
            origin__iexact=values["origin"],
            destination__iexact=values["destination"],
            vendor=vendor,
            vehicle_registration_snapshot=values["vehicle_registration"],
            vendor_freight_rate=Decimal(values["freight_amount"]),
        ).first()
    if trip:
        if trip.client_id != client.pk or trip.vendor_id != vendor.pk:
            raise ValueError(f"Trip '{trip.trip_no}' does not match the supplied client/transporter")
        return trip

    vehicle = Vehicle.objects.filter(registration_no=values["vehicle_registration"]).first()
    if vehicle and vehicle.vendor_id != vendor.pk:
        raise ValueError(f"Vehicle '{values['vehicle_registration']}' belongs to another transporter")
    if not vehicle:
        if not create_missing:
            raise ValueError(f"Vehicle '{values['vehicle_registration']}' does not exist")
        vehicle = Vehicle.objects.create(
            registration_no=values["vehicle_registration"],
            vendor=vendor,
            vehicle_type=values["vehicle_type"],
        )
    driver = Driver.objects.filter(
        name__iexact=values["driver_name"], phone=values["driver_phone"], vendor=vendor
    ).first()
    if not driver:
        if not create_missing:
            raise ValueError(f"Driver '{values['driver_name']}' does not exist")
        driver = Driver.objects.create(
            name=values["driver_name"], phone=values["driver_phone"], vendor=vendor
        )
    generated_suffix = hashlib.sha256(values["trip_key"].encode("utf-8")).hexdigest()[:10].upper()
    indent_no = values["indent_no"] or f"MIS-{job.pk}-{generated_suffix}"
    indent = Indent.objects.filter(indent_no__iexact=indent_no).first()
    if indent and indent.client_id != client.pk:
        raise ValueError(f"Indent '{indent_no}' belongs to another client")
    if not indent:
        indent = Indent.objects.create(
            indent_no=indent_no,
            client=client,
            indent_date=parse_date(values["deployment_date"]),
            origin=values["origin"],
            destination=values["destination"],
            expected_delivery_date=parse_date(values["expected_delivery_date"]),
            required_vehicle_type=values["vehicle_type"],
            customer_reference=values["record_ref"],
            notes=f"Historical MIS import #{job.pk}",
        )
    trip = Trip(
        trip_no=values["trip_no"],
        indent=indent,
        client=client,
        origin=values["origin"],
        destination=values["destination"],
        deployment_date=parse_date(values["deployment_date"]),
        expected_delivery_date=parse_date(values["expected_delivery_date"]),
        vendor=vendor,
        vehicle=vehicle,
        driver=driver,
        vendor_freight_rate=Decimal(values["freight_amount"]),
        advance_percent=Decimal(values["advance_percent"]),
        status=Trip.Status.ADVANCE_APPROVED,
        notes="\n".join(filter(None, [values["notes"], f"Historical MIS import #{job.pk}"])),
    )
    trip.save()
    _reserve_trip_number(client, trip.trip_no)
    if Decimal(values["unloading"]):
        TripCharge.objects.create(
            trip=trip,
            charge_type=TripCharge.ChargeType.UNLOADING,
            amount=Decimal(values["unloading"]),
            advance_eligible=True,
            tds_eligible=False,
            created_by=actor,
        )
    return trip


def _reserve_trip_number(client, trip_no):
    match = re.fullmatch(rf"{re.escape(client.code)}-(\d{{4}})-(\d+)", trip_no or "")
    if match:
        NumberSequence.reserve(f"trip:{client.code}:{match.group(1)}", int(match.group(2)))


def _existing_approval_item(trip, values):
    return (
        trip.approval_items.filter(
            item_status=PaymentApprovalItem.Status.APPROVED,
            gross_requested=Decimal(values["gross_approved"]),
            tds_this_request=Decimal(values["tds_approved"]),
            net_requested=Decimal(values["net_approved"]),
        )
        .order_by("-created_at")
        .first()
    )


@transaction.atomic
def confirm_mis_import(*, job, actor, create_missing=True, allow_partial=False, request_id=""):
    job = ImportJob.objects.select_for_update().get(pk=job.pk)
    if job.status == "COMPLETED":
        return job
    if getattr(actor, "role", "") != "ADMIN":
        raise PermissionError("Only an administrator can post historical financial records")
    if job.summary.get("kind") != "MIS_HISTORY":
        raise ValueError("This is not an MIS historical import")
    rows = job.summary.get("rows", [])
    invalid = [row for row in rows if not row.get("valid")]
    if invalid and not allow_partial:
        raise ValueError("The import contains invalid rows; correct them or enable valid-row-only import")
    valid_rows = [row for row in rows if row.get("valid")]
    if not valid_rows:
        raise ValueError("There are no valid MIS rows to import")

    trip_map = {}
    row_trip_map = {}
    created_trip_ids = []
    reused_trip_ids = set()
    skipped = []
    preexisting_trip_ids = set(Trip.objects.values_list("id", flat=True))
    for row in valid_rows:
        values = row["normalized"]
        key = values["trip_key"]
        if key not in trip_map:
            trip = _resolve_trip(values, actor, job, create_missing)
            trip_map[key] = trip
            if trip.pk in preexisting_trip_ids:
                reused_trip_ids.add(trip.pk)
            else:
                created_trip_ids.append(trip.pk)
        row_trip_map[row["row_no"]] = trip_map[key]

    new_item_rows = {}
    existing_items = {}
    for row in valid_rows:
        values = row["normalized"]
        trip = row_trip_map[row["row_no"]]
        if trip.pk in existing_items or trip.pk in new_item_rows:
            continue
        existing = _existing_approval_item(trip, values)
        if existing:
            existing_items[trip.pk] = existing
        else:
            new_item_rows[trip.pk] = row

    batches = {}
    for trip_id, row in new_item_rows.items():
        trip = row_trip_map[row["row_no"]]
        batch = batches.get(trip.client_id)
        if not batch:
            batch = PaymentApprovalBatch.objects.create(
                client=trip.client,
                requested_by=actor,
                submitted_at=timezone.now(),
                status=PaymentApprovalBatch.Status.APPROVED,
                purpose="HISTORICAL",
                approval_rule_snapshot={"source": "MIS_HISTORY", "import_job": job.pk},
            )
            batches[trip.client_id] = batch
        values = row["normalized"]
        item = PaymentApprovalItem.objects.create(
            batch=batch,
            trip=trip,
            vendor=trip.vendor,
            freight_rate_snapshot=Decimal(values["freight_amount"]),
            advance_percent=Decimal(values["advance_percent"]),
            freight_advance_gross=Decimal(values["calculated_advance"]),
            advance_eligible_charges=Decimal(values["unloading"]),
            advance_stage_deductions=Decimal("0"),
            gross_requested=Decimal(values["gross_approved"]),
            tds_rate=Decimal(values["tds_rate"]),
            tds_policy_snapshot=TDSPolicy.PER_PAYMENT_TAXABLE_AMOUNT,
            tds_base=Decimal(values["tds_base"]),
            tds_this_request=Decimal(values["tds_approved"]),
            net_requested=Decimal(values["net_approved"]),
            calculation_breakdown={
                "source": "MIS_HISTORY",
                "import_job": job.pk,
                "record_ref": values["record_ref"],
                "freight_rate": values["freight_amount"],
                "advance_percent": values["advance_percent"],
                "freight_advance_gross": values["calculated_advance"],
                "advance_eligible_charges": values["unloading"],
                "gross_requested": values["gross_approved"],
                "tds_rate": values["tds_rate"],
                "tds_base": values["tds_base"],
                "tds_this_payment": values["tds_approved"],
                "net_requested": values["net_approved"],
            },
            item_status=PaymentApprovalItem.Status.APPROVED,
            approver_note=f"Imported from historical MIS job #{job.pk}",
        )
        existing_items[trip_id] = item

    for batch in batches.values():
        items = batch.items.all()
        batch.gross_requested = sum((item.gross_requested for item in items), Decimal("0"))
        batch.tds_requested = sum((item.tds_this_request for item in items), Decimal("0"))
        batch.net_requested = sum((item.net_requested for item in items), Decimal("0"))
        batch.save(update_fields=["gross_requested", "tds_requested", "net_requested", "updated_at"])
        ApprovalAction.objects.create(
            actor=actor,
            action=ApprovalAction.Action.APPROVE,
            scope="BATCH",
            batch=batch,
            comment=f"Historical records approved during MIS import #{job.pk}",
            captured_totals={
                "gross": str(batch.gross_requested),
                "tds": str(batch.tds_requested),
                "net": str(batch.net_requested),
                "source": "MIS_HISTORY",
            },
        )

    payment_groups = defaultdict(list)
    seen_existing = set()
    for row in valid_rows:
        values = row["normalized"]
        if Decimal(values["gross_paid"] or "0") == 0:
            continue
        trip = row_trip_map[row["row_no"]]
        existing_payment = PaymentAllocation.objects.filter(
            trip=trip,
            payment__status=FinancePaymentTransaction.Status.PAID,
            payment__utr_reference__iexact=values["utr"],
        ).exists()
        if existing_payment:
            signature = f"{trip.pk}:{values['utr']}"
            if signature not in seen_existing:
                skipped.append({"row_no": row["row_no"], "reason": "existing_payment", "trip_id": trip.pk})
                seen_existing.add(signature)
            continue
        item = existing_items[trip.pk]
        payment_groups[(trip.vendor_id, values["payment_date"], values["utr"], values["payment_mode"])].append(
            {
                "row_no": row["row_no"],
                "item": item,
                "gross": values["gross_paid"],
                "tds": values["tds_paid"],
                "net": values["net_paid"],
            }
        )

    payment_ids = []
    for (_vendor_id, payment_date_value, utr, payment_mode), allocations in payment_groups.items():
        vendor = allocations[0]["item"].vendor
        payment = create_paid_payment(
            actor=actor,
            vendor=vendor,
            payment_date=parse_date(payment_date_value),
            utr_reference=utr,
            payment_mode=payment_mode,
            allocations=[
                {
                    "approval_item_id": entry["item"].pk,
                    "gross_amount_allocated": entry["gross"],
                    "tds_allocated": entry["tds"],
                    "net_cash_allocated": entry["net"],
                }
                for entry in allocations
            ],
            remarks=f"Historical MIS import #{job.pk}",
            request_id=request_id,
            enforce_proof=False,
            notify=False,
        )
        payment_ids.append(payment.pk)

    for row in valid_rows:
        values = row["normalized"]
        trip = row_trip_map[row["row_no"]]
        if trip.pk not in created_trip_ids:
            # Existing trips keep their live workflow state; only imported trips are stamped.
            continue
        if values["trip_status"]:
            trip.status = values["trip_status"]
        elif not trip.payment_allocations.filter(payment__status=FinancePaymentTransaction.Status.PAID).exists():
            trip.status = Trip.Status.ADVANCE_APPROVED
        trip.save(update_fields=["status", "updated_at"])

    job.status = "COMPLETED"
    job.confirmed_at = timezone.now()
    job.result = {
        "created_trip_ids": created_trip_ids,
        "reused_trip_ids": sorted(reused_trip_ids),
        "approval_ids": [batch.pk for batch in batches.values()],
        "payment_ids": payment_ids,
        "skipped": skipped + [
            {"row_no": row["row_no"], "reason": "validation_error"} for row in invalid
        ],
    }
    job.save(update_fields=["status", "confirmed_at", "result", "updated_at"])
    record_audit(
        actor=actor,
        action="MIS_HISTORY_IMPORTED",
        instance=job,
        after={
            "created_trips": len(created_trip_ids),
            "reused_trips": len(reused_trip_ids),
            "payments": len(payment_ids),
            "skipped": len(job.result["skipped"]),
        },
        request_id=request_id,
        source="IMPORT",
    )
    return job


def mis_queryset(params, user):
    queryset = Trip.objects.select_related("client", "indent", "vendor", "vehicle", "driver", "final_settlement")
    if getattr(user, "role", "") == "TRANSPORTER":
        queryset = queryset.filter(vendor_id=user.vendor_id)
    if params.get("date_from"):
        queryset = queryset.filter(deployment_date__gte=params["date_from"])
    if params.get("date_to"):
        queryset = queryset.filter(deployment_date__lte=params["date_to"])
    if params.get("vendor"):
        queryset = queryset.filter(vendor_id=params["vendor"])
    if params.get("status"):
        queryset = queryset.filter(status=params["status"])
    if params.get("q"):
        term = params["q"].strip()
        queryset = queryset.filter(
            Q(trip_no__icontains=term)
            | Q(indent__indent_no__icontains=term)
            | Q(origin__icontains=term)
            | Q(destination__icontains=term)
            | Q(vendor__display_name__icontains=term)
            | Q(vehicle_registration_snapshot__icontains=term)
            | Q(payment_allocations__payment__utr_reference__icontains=term)
        ).distinct()
    return queryset.order_by("-deployment_date", "-id").prefetch_related(
        Prefetch(
            "approval_items",
            queryset=PaymentApprovalItem.objects.filter(
                item_status=PaymentApprovalItem.Status.APPROVED
            ).select_related("batch"),
            to_attr="mis_approved_items",
        ),
        Prefetch(
            "payment_allocations",
            queryset=PaymentAllocation.objects.filter(
                payment__status=FinancePaymentTransaction.Status.PAID
            ).select_related("payment"),
            to_attr="mis_paid_allocations",
        ),
    )


def mis_row(trip):
    approved_gross = sum((item.gross_requested for item in trip.mis_approved_items), Decimal("0"))
    cash_paid = sum((item.net_cash_allocated for item in trip.mis_paid_allocations), Decimal("0"))
    tds_paid = sum((item.tds_allocated for item in trip.mis_paid_allocations), Decimal("0"))
    gross_paid = cash_paid + tds_paid
    liability = getattr(trip, "final_settlement", None)
    liability = liability.total_vendor_gross_cost if liability else trip.vendor_freight_rate
    approved_outstanding = max(Decimal("0"), approved_gross - gross_paid)
    if not gross_paid:
        payment_status = "UNPAID"
    elif approved_outstanding:
        payment_status = "PARTIAL"
    else:
        payment_status = "PAID"
    payment_map = {}
    for allocation in trip.mis_paid_allocations:
        payment = allocation.payment
        entry = payment_map.setdefault(
            payment.pk,
            {
                "id": payment.pk,
                "payment_no": payment.payment_no,
                "payment_date": payment.payment_date,
                "utr": payment.utr_reference,
                "gross": Decimal("0"),
                "tds": Decimal("0"),
                "net": Decimal("0"),
            },
        )
        entry["gross"] += allocation.gross_amount_allocated
        entry["tds"] += allocation.tds_allocated
        entry["net"] += allocation.net_cash_allocated
    return {
        "id": trip.pk,
        "trip_no": trip.trip_no,
        "indent_no": trip.indent.indent_no,
        "client": trip.client.name,
        "deployment_date": trip.deployment_date,
        "created_at": trip.created_at,
        "route": f"{trip.origin} → {trip.destination}",
        "vendor_id": trip.vendor_id,
        "vendor": trip.vendor.display_name,
        "vehicle_no": trip.vehicle_registration_snapshot,
        "driver": trip.driver_name_snapshot,
        "freight_100": trip.vendor_freight_rate,
        "advance_percent": trip.advance_percent,
        "approved_gross": approved_gross,
        "cash_paid": cash_paid,
        "tds_paid": tds_paid,
        "gross_accounted": gross_paid,
        "approved_outstanding": approved_outstanding,
        "total_remaining": max(Decimal("0"), liability - gross_paid),
        "trip_status": trip.status,
        "payment_status": payment_status,
        "payments": list(payment_map.values()),
    }


def filter_mis_rows(rows, payment_status):
    return [row for row in rows if not payment_status or row["payment_status"] == payment_status]


def mis_summary(rows):
    fields = (
        "freight_100",
        "approved_gross",
        "cash_paid",
        "tds_paid",
        "gross_accounted",
        "approved_outstanding",
        "total_remaining",
    )
    return {
        "record_count": len(rows),
        **{
            field: sum((Decimal(str(row[field])) for row in rows), Decimal("0"))
            for field in fields
        },
    }


def export_mis_rows(rows):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "MIS Register"
    columns = [
        "Trip", "Indent", "Client", "Deployment date", "Route", "Transporter", "Vehicle",
        "Driver", "Freight 100%", "Approved gross", "Cash paid", "TDS paid",
        "Gross accounted", "Approved outstanding", "Total remaining", "Trip status",
        "Payment status", "Payment references",
    ]
    sheet.append(columns)
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="D71920")
    for row in rows:
        sheet.append(
            [
                row["trip_no"], row["indent_no"], row["client"], row["deployment_date"], row["route"],
                row["vendor"], row["vehicle_no"], row["driver"], float(row["freight_100"]),
                float(row["approved_gross"]), float(row["cash_paid"]), float(row["tds_paid"]),
                float(row["gross_accounted"]), float(row["approved_outstanding"]),
                float(row["total_remaining"]), row["trip_status"], row["payment_status"],
                "; ".join(f"{payment['payment_no']} / {payment['utr']}" for payment in row["payments"]),
            ]
        )
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()
