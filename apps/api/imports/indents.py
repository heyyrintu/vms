import hashlib
from datetime import date, datetime, time
from decimal import Decimal, InvalidOperation
from io import BytesIO

from django.db import transaction
from django.utils import timezone
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.utils.datetime import from_excel

from audit.models import record_audit
from operations.models import Client, Indent

from .models import ImportJob

INDENT_COLUMNS = [
    "FROM",
    "CHALLAN NO",
    "CHALLAN DATE & TIME",
    "SHIP TO PARTY CODE",
    "SHIP TO PARTY NAME",
    "ADDRESS",
    "TO LOCATION",
    "TO STATE",
    "PIN CODE",
    "ITEM",
    "DEF QTY",
    "QTY LTR",
    "REPORTING DATE & TIME",
    "EDD",
    "BRANCH",
    "COST CENTER",
    "NOTES",
]

REQUIRED_INDENT_COLUMNS = INDENT_COLUMNS[:12]


def _text(value):
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _decimal(value, field, errors):
    if value in (None, ""):
        errors.append(f"{field} is required")
        return None
    try:
        result = Decimal(str(value)).quantize(Decimal("0.001"))
    except (InvalidOperation, ValueError):
        errors.append(f"{field} must be numeric")
        return None
    if result < 0:
        errors.append(f"{field} cannot be negative")
    return result


def _datetime(value, field, errors, epoch, required=False):
    if value in (None, ""):
        if required:
            errors.append(f"{field} is required")
        return None
    parsed = None
    if isinstance(value, datetime):
        parsed = value
    elif isinstance(value, date):
        parsed = datetime.combine(value, time.min)
    elif isinstance(value, int | float):
        try:
            converted = from_excel(value, epoch)
            parsed = converted if isinstance(converted, datetime) else datetime.combine(converted, time.min)
        except (ValueError, OverflowError):
            pass
    elif isinstance(value, str):
        for fmt in (
            "%Y-%m-%d %H:%M",
            "%Y-%m-%d %H:%M:%S",
            "%d/%m/%Y %H:%M",
            "%d-%m-%Y %H:%M",
            "%Y-%m-%d",
            "%d/%m/%Y",
            "%d-%m-%Y",
        ):
            try:
                parsed = datetime.strptime(value.strip(), fmt)
                break
            except ValueError:
                continue
    if parsed is None:
        errors.append(f"{field} is not a valid date/time")
        return None
    if timezone.is_naive(parsed):
        parsed = timezone.make_aware(parsed, timezone.get_current_timezone())
    return parsed


def preview_indent_workbook(content, *, client):
    workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
    sheet = workbook["Indent Upload"] if "Indent Upload" in workbook.sheetnames else workbook.active
    headers = [_text(cell.value).upper() for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
    missing = [column for column in REQUIRED_INDENT_COLUMNS if column not in headers]
    if missing:
        raise ValueError(f"Missing required columns: {', '.join(missing)}")
    positions = {name: headers.index(name) for name in INDENT_COLUMNS if name in headers}
    rows = []
    seen = set()
    for row_no, values in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
        if not any(value not in (None, "") for value in values):
            continue
        if len(rows) >= 2000:
            raise ValueError("Indent imports are limited to 2,000 populated rows per workbook")

        def get(name, row_values=values):
            position = positions.get(name)
            return row_values[position] if position is not None and position < len(row_values) else None

        errors, warnings = [], []
        challan_no = _text(get("CHALLAN NO")).upper()
        required_text = {
            "FROM": _text(get("FROM")),
            "CHALLAN NO": challan_no,
            "SHIP TO PARTY CODE": _text(get("SHIP TO PARTY CODE")),
            "SHIP TO PARTY NAME": _text(get("SHIP TO PARTY NAME")),
            "ADDRESS": _text(get("ADDRESS")),
            "TO LOCATION": _text(get("TO LOCATION")),
            "TO STATE": _text(get("TO STATE")),
            "PIN CODE": _text(get("PIN CODE")),
            "ITEM": _text(get("ITEM")),
        }
        for field, value in required_text.items():
            if not value:
                errors.append(f"{field} is required")
        if required_text["PIN CODE"] and not required_text["PIN CODE"].isdigit():
            errors.append("PIN CODE must contain digits only")
        challan_at = _datetime(get("CHALLAN DATE & TIME"), "CHALLAN DATE & TIME", errors, workbook.epoch, required=True)
        reporting_at = _datetime(get("REPORTING DATE & TIME"), "REPORTING DATE & TIME", errors, workbook.epoch)
        edd_at = _datetime(get("EDD"), "EDD", errors, workbook.epoch)
        default_quantity = _decimal(get("DEF QTY"), "DEF QTY", errors)
        quantity_ltrs = _decimal(get("QTY LTR"), "QTY LTR", errors)
        if challan_no in seen:
            errors.append("Duplicate CHALLAN NO in this workbook")
        seen.add(challan_no)
        if challan_no and Indent.objects.filter(client=client, challan_no__iexact=challan_no).exists():
            warnings.append("This challan already exists and will be skipped")
        rows.append(
            {
                "row_no": row_no,
                "valid": not errors,
                "errors": errors,
                "warnings": warnings,
                "normalized": {
                    "origin": required_text["FROM"],
                    "challan_no": challan_no,
                    "challan_datetime": challan_at.isoformat() if challan_at else None,
                    "ship_to_party_code": required_text["SHIP TO PARTY CODE"],
                    "ship_to_party_name": required_text["SHIP TO PARTY NAME"],
                    "ship_to_address": required_text["ADDRESS"],
                    "destination": required_text["TO LOCATION"],
                    "destination_state": required_text["TO STATE"],
                    "pin_code": required_text["PIN CODE"],
                    "item": required_text["ITEM"],
                    "default_quantity": str(default_quantity) if default_quantity is not None else None,
                    "quantity_ltrs": str(quantity_ltrs) if quantity_ltrs is not None else None,
                    "reporting_datetime": reporting_at.isoformat() if reporting_at else None,
                    "expected_delivery_date": edd_at.date().isoformat() if edd_at else None,
                    "branch": _text(get("BRANCH")),
                    "cost_center": _text(get("COST CENTER")),
                    "notes": _text(get("NOTES")),
                },
            }
        )
    if not rows:
        raise ValueError("The workbook does not contain any indent rows")
    source_hash = hashlib.sha256(content + f":indent:{client.pk}".encode()).hexdigest()
    return {
        "source_hash": source_hash,
        "columns": headers,
        "required_columns": REQUIRED_INDENT_COLUMNS,
        "row_count": len(rows),
        "valid_count": sum(row["valid"] for row in rows),
        "error_count": sum(not row["valid"] for row in rows),
        "duplicate_count": sum(bool(row["warnings"]) for row in rows),
        "rows": rows,
    }


def _canonical_indent_no(client, challan_no):
    candidate = challan_no[:50]
    if not Indent.objects.filter(indent_no=candidate).exists():
        return candidate
    base = f"{client.code}-{challan_no}"[:45]
    candidate = base
    number = 1
    while Indent.objects.filter(indent_no=candidate).exists():
        number += 1
        candidate = f"{base[:44]}-{number}"[:50]
    return candidate


@transaction.atomic
def confirm_indent_import(*, job, actor, client_id, allow_partial=False, request_id=""):
    job = ImportJob.objects.select_for_update().get(pk=job.pk)
    if job.status == "COMPLETED":
        return job
    if job.uploaded_by_id != actor.pk and getattr(actor, "role", "") != "ADMIN":
        raise PermissionError("Only the uploader or an administrator can confirm this import")
    client = Client.objects.get(pk=client_id)
    rows = job.summary.get("rows", [])
    invalid = [row for row in rows if not row.get("valid")]
    if invalid and not allow_partial:
        raise ValueError("The import contains invalid rows; correct them or enable valid-row-only import")
    created, skipped = [], []
    for row in rows:
        if not row.get("valid"):
            skipped.append({"row_no": row["row_no"], "reason": "validation_error"})
            continue
        values = row["normalized"]
        duplicate = Indent.objects.filter(client=client, challan_no__iexact=values["challan_no"]).first()
        if duplicate:
            skipped.append({"row_no": row["row_no"], "reason": "duplicate_challan", "indent_id": duplicate.pk})
            continue
        challan_at = datetime.fromisoformat(values["challan_datetime"])
        indent = Indent.objects.create(
            indent_no=_canonical_indent_no(client, values["challan_no"]),
            challan_no=values["challan_no"],
            challan_datetime=challan_at,
            indent_date=challan_at.date(),
            client=client,
            origin=values["origin"],
            destination=values["destination"],
            ship_to_party_code=values["ship_to_party_code"],
            ship_to_party_name=values["ship_to_party_name"],
            ship_to_address=values["ship_to_address"],
            destination_state=values["destination_state"],
            pin_code=values["pin_code"],
            item=values["item"],
            default_quantity=values["default_quantity"],
            quantity_ltrs=values["quantity_ltrs"],
            quantity=values["default_quantity"],
            total_load=values["quantity_ltrs"],
            reporting_datetime=values["reporting_datetime"],
            expected_delivery_date=values["expected_delivery_date"],
            branch=values["branch"],
            cost_center=values["cost_center"],
            notes=values["notes"],
            uom_ltrs="LTR",
            status=Indent.Status.OPEN,
        )
        record_audit(
            actor=actor,
            action="INDENT_IMPORTED",
            instance=indent,
            after={"import_job": job.pk, "source_row": row["row_no"], "challan_no": indent.challan_no},
            request_id=request_id,
            source="IMPORT",
        )
        created.append(indent.pk)
    job.status = "COMPLETED"
    job.confirmed_at = timezone.now()
    job.result = {"created_indent_ids": created, "skipped": skipped, "errors": []}
    job.save(update_fields=["status", "confirmed_at", "result", "updated_at"])
    record_audit(
        actor=actor,
        action="INDENT_IMPORT_CONFIRMED",
        instance=job,
        after={"created": len(created), "skipped": len(skipped)},
        request_id=request_id,
        source="IMPORT",
    )
    return job


def build_indent_template():
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Indent Upload"
    sheet.append(INDENT_COLUMNS)
    sheet.append(
        [
            "Sonipat",
            "8127959325",
            datetime(2026, 9, 3, 10, 30),
            "2088979",
            "SAGAR MOTOR",
            "BY PASS ROAD KICHHA UDHAM SINGH NAGAR, KICHHA ROAD, 263153, INDIA",
            "Kichha",
            "Uttarakhand",
            "263153",
            "20 LTR",
            Decimal("300"),
            Decimal("6000"),
            datetime(2026, 9, 3, 12, 0),
            datetime(2026, 9, 4, 18, 0),
            "Sonipat",
            "",
            "Delete or replace this sample row before upload.",
        ]
    )
    header_fill = PatternFill("solid", fgColor="17365D")
    required_fill = PatternFill("solid", fgColor="FFF2CC")
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for cell in sheet[2][:12]:
        cell.fill = required_fill
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(INDENT_COLUMNS))}2"
    widths = [16, 18, 22, 20, 26, 48, 20, 18, 12, 18, 14, 14, 24, 22, 18, 18, 42]
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.row_dimensions[1].height = 34
    sheet.row_dimensions[2].height = 70
    for row in sheet.iter_rows(min_row=2, max_row=2001):
        row[2].number_format = "yyyy-mm-dd hh:mm"
        row[12].number_format = "yyyy-mm-dd hh:mm"
        row[13].number_format = "yyyy-mm-dd hh:mm"
        row[10].number_format = "#,##0.000"
        row[11].number_format = "#,##0.000"
    notes = workbook.create_sheet("Instructions")
    notes.append(["Indent upload instructions"])
    notes.append(["1. Keep the column names unchanged."])
    notes.append(["2. Enter one indent/challan per row. Required columns are the first 12 columns."])
    notes.append(["3. Use yyyy-mm-dd hh:mm for date/time values. Date-only values are accepted as midnight."])
    notes.append(["4. Keep challan numbers, ship-to codes and PIN codes as text when they contain leading zeroes."])
    notes.append(["5. Remove the yellow sample row before uploading your production records."])
    notes.column_dimensions["A"].width = 105
    notes["A1"].font = Font(size=14, bold=True, color="17365D")
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()
