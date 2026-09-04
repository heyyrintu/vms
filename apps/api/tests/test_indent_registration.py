from io import BytesIO

import pytest
from django.core.files.uploadedfile import SimpleUploadedFile
from openpyxl import Workbook, load_workbook
from rest_framework.test import APIClient

from accounts.models import User
from operations.models import Client, Driver, Indent, Trip, TripIndent, Vehicle, Vendor


@pytest.mark.django_db
def test_indent_registration_captures_challan_fields_and_timestamp(users):
    client = Client.objects.create(code="NPL", name="NPL")
    api = APIClient()
    api.force_authenticate(users[User.Role.OPERATIONS])

    response = api.post(
        "/api/indents/",
        {
            "client": client.pk,
            "challan_no": "8127959325",
            "challan_datetime": "2026-09-03T10:30:00+05:30",
            "origin": "Sonipat",
            "ship_to_party_code": "2088979",
            "ship_to_party_name": "Sagar Motor",
            "ship_to_address": "By Pass Road, Kichha",
            "destination": "Kichha",
            "destination_state": "Uttarakhand",
            "pin_code": "263153",
            "item": "20 LTR",
            "default_quantity": "300.000",
            "quantity_ltrs": "6000.000",
            "status": "OPEN",
        },
        format="json",
    )

    assert response.status_code == 201, response.data
    indent = Indent.objects.get(pk=response.data["id"])
    assert indent.indent_no == "8127959325"
    assert indent.indent_date.isoformat() == "2026-09-03"
    assert indent.quantity_ltrs == 6000
    assert response.data["trip_count"] == 0


@pytest.mark.django_db
def test_trip_accepts_multiple_indents_and_suggests_route_matches(users):
    client = Client.objects.create(code="NPL", name="NPL")
    vendor = Vendor.objects.create(vendor_code="V001", legal_name="Vendor", display_name="Vendor")
    vehicle = Vehicle.objects.create(registration_no="HR10AB1234", vendor=vendor, vehicle_type="32 FT")
    driver = Driver.objects.create(name="Driver", phone="919800000001", vendor=vendor)
    indents = [
        Indent.objects.create(
            indent_no=f"CH-{index}", challan_no=f"CH-{index}", client=client,
            indent_date="2026-09-03", origin="Sonipat", destination="Kichha",
        )
        for index in range(1, 4)
    ]
    api = APIClient()
    api.force_authenticate(users[User.Role.OPERATIONS])

    response = api.post(
        "/api/trips/",
        {
            "indent": indents[0].pk,
            "indent_ids": [indents[0].pk, indents[1].pk],
            "client": client.pk,
            "origin": "Sonipat",
            "destination": "Kichha",
            "deployment_date": "2026-09-03",
            "vendor": vendor.pk,
            "vehicle": vehicle.pk,
            "driver": driver.pk,
            "vendor_freight_rate": "50000.00",
        },
        format="json",
    )

    assert response.status_code == 201, response.data
    trip = Trip.objects.get(pk=response.data["id"])
    assert list(trip.indents.values_list("pk", flat=True)) == [indents[0].pk, indents[1].pk]
    assert TripIndent.objects.get(trip=trip, indent=indents[0]).is_primary is True
    assert response.data["indent_count"] == 2

    suggestions = api.get(f"/api/trips/{trip.pk}/indent-suggestions/")
    assert suggestions.status_code == 200
    assert suggestions.data[0]["id"] == indents[2].pk
    assert suggestions.data[0]["match_score"] == 8


@pytest.mark.django_db
def test_indent_excel_preview_confirm_and_template(users):
    client = Client.objects.create(code="NPL", name="NPL")
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Indent Upload"
    sheet.append([
        "FROM", "CHALLAN NO", "CHALLAN DATE & TIME", "SHIP TO PARTY CODE",
        "SHIP TO PARTY NAME", "ADDRESS", "TO LOCATION", "TO STATE", "PIN CODE",
        "ITEM", "DEF QTY", "QTY LTR",
    ])
    sheet.append([
        "Sonipat", "8127959325", "2026-09-03 10:30", "2088979", "Sagar Motor",
        "By Pass Road", "Kichha", "Uttarakhand", "263153", "20 LTR", 300, 6000,
    ])
    content = BytesIO()
    workbook.save(content)
    api = APIClient()
    api.force_authenticate(users[User.Role.OPERATIONS])

    preview = api.post(
        "/api/imports/indents/preview/",
        {
            "client": str(client.pk),
            "file": SimpleUploadedFile(
                "indents.xlsx",
                content.getvalue(),
                content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            ),
        },
        format="multipart",
    )
    assert preview.status_code == 200, preview.data
    assert preview.data["valid_count"] == 1
    assert preview.data["rows"][0]["normalized"]["quantity_ltrs"] == "6000.000"

    confirmed = api.post(
        f"/api/imports/indents/{preview.data['job_id']}/confirm/",
        {"allow_partial": False},
        format="json",
    )
    assert confirmed.status_code == 200, confirmed.data
    assert len(confirmed.data["result"]["created_indent_ids"]) == 1
    assert Indent.objects.get(challan_no="8127959325").ship_to_party_name == "Sagar Motor"

    template = api.get("/api/imports/indents/template/")
    assert template.status_code == 200
    template_book = load_workbook(BytesIO(template.content), read_only=True)
    assert template_book.sheetnames == ["Indent Upload", "Instructions"]
    assert template_book["Indent Upload"]["C1"].value == "CHALLAN DATE & TIME"
