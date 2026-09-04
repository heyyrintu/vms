from io import BytesIO

import pytest
from django.conf import settings
from openpyxl import load_workbook
from rest_framework.test import APIClient

from accounts.models import User
from approvals.services import create_approval_batch


@pytest.mark.django_db
def test_transporter_cannot_read_another_vendor_trip(trip_factory, users):
    own_trip = trip_factory(1)
    other_trip = trip_factory(2)
    transporter = User.objects.create_user(
        username="vendor-user", password="StrongPass123!", role=User.Role.TRANSPORTER, vendor=own_trip.vendor
    )
    client = APIClient()
    client.force_authenticate(transporter)
    assert client.get(f"/api/trips/{own_trip.pk}/").status_code == 200
    assert client.get(f"/api/trips/{other_trip.pk}/").status_code == 404
    assert client.get(f"/api/vendor-ledger/{other_trip.vendor_id}/").status_code == 403


@pytest.mark.django_db
def test_operations_cannot_approve_or_pay(trip_factory, users):
    trip = trip_factory(1)
    client = APIClient()
    client.force_authenticate(users[User.Role.OPERATIONS])
    created = client.post("/api/approval-batches/", {"trip_ids": [trip.pk]}, format="json")
    assert created.status_code == 201
    batch_id = created.data["id"]
    assert client.post(f"/api/approval-batches/{batch_id}/submit/", {}, format="json").status_code == 200
    assert client.post(f"/api/approval-batches/{batch_id}/decide/", {"decision": "APPROVE"}, format="json").status_code == 403


@pytest.mark.django_db
def test_session_login_uses_httponly_cookie(users):
    client = APIClient(enforce_csrf_checks=True)
    response = client.get("/api/auth/csrf/")
    token = response.data["csrfToken"]
    login = client.post(
        "/api/auth/login/",
        {"username": "operations", "password": "StrongPass123!"},
        format="json",
        HTTP_X_CSRFTOKEN=token,
    )
    assert login.status_code == 200
    assert client.cookies["sessionid"]["httponly"] is True


@pytest.mark.django_db
def test_session_requests_accept_next_development_fallback_port(users):
    origin = "http://localhost:3002"
    assert origin in settings.CSRF_TRUSTED_ORIGINS
    assert origin in settings.CORS_ALLOWED_ORIGINS

    client = APIClient(enforce_csrf_checks=True)
    csrf = client.get("/api/auth/csrf/", HTTP_ORIGIN=origin)
    login = client.post(
        "/api/auth/login/",
        {"username": "operations", "password": "StrongPass123!"},
        format="json",
        HTTP_ORIGIN=origin,
        HTTP_X_CSRFTOKEN=csrf.data["csrfToken"],
    )
    assert login.status_code == 200

    logout = client.post(
        "/api/auth/logout/",
        {},
        format="json",
        HTTP_ORIGIN=origin,
        HTTP_X_CSRFTOKEN=client.cookies["csrftoken"].value,
    )
    assert logout.status_code == 204


@pytest.mark.django_db
def test_transporter_secondary_endpoints_are_vendor_isolated(trip_factory):
    own_trip = trip_factory(1)
    other_trip = trip_factory(2)
    transporter = User.objects.create_user(
        username="isolated-vendor",
        password="StrongPass123!",
        role=User.Role.TRANSPORTER,
        vendor=own_trip.vendor,
    )
    client = APIClient()
    client.force_authenticate(transporter)

    assert [row["id"] for row in client.get("/api/vehicles/").data["results"]] == [
        own_trip.vehicle_id
    ]
    assert [row["id"] for row in client.get("/api/drivers/").data["results"]] == [
        own_trip.driver_id
    ]
    assert client.get("/api/finance/pending/").status_code == 403
    assert client.get("/api/billings/").data["count"] == 0

    export = client.get("/api/imports/excel/export/")
    workbook = load_workbook(BytesIO(export.content), read_only=True)
    rows = list(workbook.active.iter_rows(values_only=True))
    assert len(rows) == 2
    assert rows[1][1:3] == (own_trip.origin, own_trip.destination)
    assert other_trip.trip_no not in str(rows)


@pytest.mark.django_db
def test_transporter_cannot_open_internal_multi_vendor_approval(trip_factory, users):
    own_trip = trip_factory(1)
    other_trip = trip_factory(2)
    batch = create_approval_batch(
        actor=users[User.Role.OPERATIONS], trips=[own_trip, other_trip]
    )
    transporter = User.objects.create_user(
        username="approval-vendor",
        password="StrongPass123!",
        role=User.Role.TRANSPORTER,
        vendor=own_trip.vendor,
    )
    client = APIClient()
    client.force_authenticate(transporter)

    response = client.get(f"/api/approval-batches/{batch.pk}/")
    assert response.status_code == 403
