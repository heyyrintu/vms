from io import BytesIO

import pytest
from openpyxl import Workbook, load_workbook
from rest_framework.test import APIClient

from accounts.models import User
from approvals.models import ApprovalAction, PaymentApprovalItem
from imports.mis import MIS_COLUMNS, build_mis_template, preview_mis_workbook
from integrations.models import IntegrationMessage
from operations.models import Client, Trip
from payments.models import FinancePaymentTransaction, PaymentAllocation, TDSEntry


def mis_workbook(*rows):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "MIS Records"
    sheet.append(MIS_COLUMNS)
    for values in rows:
        sheet.append([values.get(column, "") for column in MIS_COLUMNS])
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def historical_row(**overrides):
    row = {
        "RECORD REF": "OLD-0001",
        "TRIP NO": "NPL-2024-000001",
        "CLIENT CODE": "NPL",
        "INDENT NO": "IND-2024-001",
        "FROM": "Sonipat",
        "TO": "Delhi",
        "DEPLOYMENT DATE": "2024-07-27",
        "EDD": "2024-07-28",
        "TRANSPORTER CODE": "V-OLD-001",
        "TRANSPORTER NAME": "Historical Transporter",
        "VEHICLE NO": "HR 10 AB 1234",
        "VEHICLE TYPE": "32 FT",
        "DRIVER NAME": "Historical Driver",
        "DRIVER NO": "9000000000",
        "FREIGHT AMOUNT": "50000",
        "ADVANCE PERCENT": "90",
        "UNLOADING": "2500",
        "GROSS APPROVED": "47500",
        "TDS RATE": "1",
        "TDS BASE": "45000",
        "TDS APPROVED": "450",
        "NET APPROVED": "47050",
        "TDS PAID": "450",
        "NET PAID": "47050",
        "PAYMENT DATE": "2024-07-27",
        "UTR": "OLD-UTR-001",
        "PAYMENT MODE": "BANK_TRANSFER",
        "TRIP STATUS": "DELIVERED",
        "NOTES": "Opening MIS record",
    }
    row.update(overrides)
    return row


def test_mis_template_has_blank_upload_sheet_and_separate_example():
    workbook = load_workbook(BytesIO(build_mis_template()), data_only=True)
    assert workbook["MIS Records"].max_row == 1
    assert workbook["Example"].max_row == 2
    assert [cell.value for cell in workbook["MIS Records"][1]] == MIS_COLUMNS


@pytest.mark.django_db
def test_mis_preview_validates_reconciliation(organization):
    content = mis_workbook(historical_row(**{"NET APPROVED": "47000"}))
    preview = preview_mis_workbook(content)
    assert preview["error_count"] == 1
    assert "NET APPROVED must equal" in " ".join(preview["rows"][0]["errors"])


@pytest.mark.django_db
def test_admin_can_import_and_track_historical_trip_payment(users, organization, tmp_path, settings):
    settings.MEDIA_ROOT = tmp_path
    Client.objects.create(code="NPL", name="NPL")
    client = APIClient()
    client.force_authenticate(users[User.Role.ADMIN])
    content = mis_workbook(historical_row())
    upload = BytesIO(content)
    upload.name = "mis-history.xlsx"

    preview = client.post("/api/mis/import/preview/", {"file": upload}, format="multipart")
    assert preview.status_code == 200, preview.data
    assert preview.data["valid_count"] == 1

    confirmed = client.post(
        f"/api/mis/import/{preview.data['job_id']}/confirm/",
        {"create_missing": True},
        format="json",
    )
    assert confirmed.status_code == 200, confirmed.data
    result = confirmed.data["result"]
    assert len(result["created_trip_ids"]) == 1
    assert len(result["approval_ids"]) == 1
    assert len(result["payment_ids"]) == 1

    trip = Trip.objects.get(pk=result["created_trip_ids"][0])
    assert trip.trip_no == "NPL-2024-000001"
    assert trip.status == Trip.Status.DELIVERED
    item = PaymentApprovalItem.objects.get(trip=trip)
    assert item.item_status == PaymentApprovalItem.Status.APPROVED
    payment = FinancePaymentTransaction.objects.get(pk=result["payment_ids"][0])
    assert payment.status == FinancePaymentTransaction.Status.PAID
    assert str(payment.net_paid_amount) == "47050.00"
    assert PaymentAllocation.objects.filter(payment=payment, trip=trip).count() == 1
    assert TDSEntry.objects.filter(payment=payment, trip=trip, tds_amount="450.00").count() == 1
    assert ApprovalAction.objects.filter(batch=item.batch, action=ApprovalAction.Action.APPROVE).exists()
    assert IntegrationMessage.objects.count() == 0

    register = client.get("/api/mis/records/?q=OLD-UTR-001&payment_status=PAID")
    assert register.status_code == 200
    assert register.data["count"] == 1
    assert register.data["rows"][0]["trip_no"] == trip.trip_no
    assert str(register.data["summary"]["cash_paid"]) == "47050.00"
    assert str(register.data["summary"]["tds_paid"]) == "450.00"
    assert str(register.data["summary"]["total_remaining"]) == "2500.00"

    repeated = client.post(
        f"/api/mis/import/{preview.data['job_id']}/confirm/",
        {"create_missing": True},
        format="json",
    )
    assert repeated.status_code == 200
    assert FinancePaymentTransaction.objects.count() == 1


@pytest.mark.django_db
def test_non_admin_cannot_preview_historical_financial_import(users, organization):
    client = APIClient()
    client.force_authenticate(users[User.Role.FINANCE])
    upload = BytesIO(mis_workbook(historical_row()))
    upload.name = "mis-history.xlsx"
    response = client.post("/api/mis/import/preview/", {"file": upload}, format="multipart")
    assert response.status_code == 403
